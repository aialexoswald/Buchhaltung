from datetime import datetime

import openpyxl
import pytest

from buchhaltung.importer import importieren


@pytest.fixture
def excel(tmp_path):
    """Kleine Excel-Datei im Aufbau der bisherigen Buchhaltungsdatei."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ar = wb.create_sheet("AR_2026")
    ar.append(["ReNr", "Empfänger", "Datum", "netto ", "brutto", "bezahlt", "MwSt", None, None, "Thema/ Projekt"])
    ar.append(["2026_001", "ETC ", datetime(2026, 1, 10), 1042.5, 1251, 0, 208.5])
    ar.append(["2026_002", "Wifi Werbeakademie", datetime(2026, 1, 18), 1000, 1000, 1, 0])
    ar.append(["2026_003", "Innung", datetime(2026, 8, 10), 2990, 3837.18, 1, 598, None, "inkl. Zug", None])
    ar.append(["2026_004", "POSLOVNI MEDIJI", datetime(2026, 8, 21), 3680, 3680, 1, 0, None, "Land Kärnten", "reverse charge"])
    ar.append(["2026_005", "APA", datetime(2027, 5, 4), 4550, 5460, 0, 910])
    ar.append(["2026_006", None, datetime(2026, 1, 10), None, 0, 1, 0])  # leere Vorlagezeile
    ar.append([None, None, None, None, None, None, None, 1042.5, "Jänner"])  # Monatssumme

    alt = wb.create_sheet("AR_2020")
    alt.append([None, "AR FUTURA GmbH"])
    alt.append([None, "ReNr", "Empfänger", "Datum", "netto ", "brutto", "bezahlt", "MwSt"])
    alt.append([None, "2020_001", "Penny", datetime(2020, 1, 2), 9960, 11952, 0, 1992])

    er = wb.create_sheet("ER-2026")
    er.append([None, None, None, "Eingangsrechnungsbuch"])
    er.append([None, None, None, "Futura GmbH"])
    er.append(["ENR", "Datum", "LNR", "Lieferant", "Währung", "Betrag", "MwSt.-Satz", "Text (allgemein)", "Text (Zusatz)",
               "Kosten-art", "Netto", "MwSt", "USt.-Code", "Sonder-code", "UID", "Kostenstelle", "weiter-verrechnen an",
               "verrechnet", "überwiesen"])
    zeile = 4
    for monat in range(1, 13):
        start = zeile
        if monat == 1:
            er.append([1, datetime(2024, 1, 1), None, "Spinnwerk", None, 450, None, None, "Webhosting", "#REF!", 375, 75,
                       None, None, None, None, "Futura", "intern", "ja"])
            er.append([2, datetime(2026, 1, 20), None, "Druck.at", None, 96, None, None, "RollUp", None, 80, 16,
                       None, None, None, None, "EG Tullnerbach", "nein", None])
            er.append([3, None, None, "Uniqa", None, 563.34, None, None, "KFZ-Versicherung", None, 563.34, 0,
                       None, None, None, None, "Futura", "intern", "???"])
            zeile += 3
        else:
            er.append([None, None, None, None, None, None])
            zeile += 1
        er.append([None] * 10 + [f"=SUM(K{start}:K{zeile - 1})"])
        zeile += 1
    er.append([None] * 10 + ["=SUM(K4:K60)"])  # Gesamtsumme – muss ignoriert werden

    pfad = tmp_path / "AR_ER.xlsx"
    wb.save(pfad)
    return pfad


def test_import_ar(con, excel):
    bericht = importieren(con, excel)
    assert bericht.ar == 6
    r = {z["nummer"]: z for z in con.execute(
        "SELECT ar.*, k.name AS kunde FROM ausgangsrechnungen ar JOIN kunden k ON k.id = ar.kunde_id")}
    assert set(r) == {"2026_001", "2026_002", "2026_003", "2026_004", "2026_005", "2020_001"}
    assert (r["2026_001"]["netto_cent"], r["2026_001"]["brutto_cent"], r["2026_001"]["kunde"]) == (104250, 125100, "ETC")
    assert r["2026_002"]["ust_cent"] == 0
    # Barauslagen ohne USt werden als eigene Position erkannt, Brutto bleibt erhalten
    assert r["2026_003"]["brutto_cent"] == 383718
    # „bezahlt“ = 1 bedeutet in der bisherigen Datei: noch offen
    assert r["2026_001"]["bezahlt"] == 1 and r["2026_002"]["bezahlt"] == 0
    assert r["2026_004"]["steuerhinweis"].startswith("Reverse Charge")
    assert r["2026_004"]["projekt"] == "reverse charge"
    assert any("2026_005" in w for w in bericht.warnungen)


def test_import_er(con, excel):
    importieren(con, excel)
    er = {z["lieferant"]: z for z in con.execute("SELECT * FROM eingangsrechnungen")}
    assert len(er) == 3
    # Datum 2024 im Jänner-Block → geschätzt auf Jänner 2026
    assert (er["Spinnwerk"]["datum"], er["Spinnwerk"]["datum_geschaetzt"]) == ("2026-01-01", 1)
    assert (er["Druck.at"]["datum"], er["Druck.at"]["datum_geschaetzt"]) == ("2026-01-20", 0)
    assert er["Spinnwerk"]["weiter_kunde_id"] is None and er["Spinnwerk"]["kostenart"] == ""
    assert er["Druck.at"]["weiter_kunde_id"] is not None and er["Druck.at"]["weiter_erledigt"] == 0
    # Versicherung ohne MwSt → 0 %
    assert (er["Uniqa"]["ust_satz"], er["Uniqa"]["vorsteuer_cent"], er["Uniqa"]["bezahlt"]) == (0, 0, 0)
    assert er["Spinnwerk"]["vorsteuer_cent"] == 7500


def test_import_wiederholbar(con, excel):
    importieren(con, excel)
    zweiter = importieren(con, excel)
    assert zweiter.ar == 0 and zweiter.er == 0
    assert con.execute("SELECT COUNT(*) FROM ausgangsrechnungen").fetchone()[0] == 6


def test_offen_wert_umkehrbar(con, excel):
    importieren(con, excel, offen_wert=0)
    bezahlt = con.execute("SELECT bezahlt FROM ausgangsrechnungen WHERE nummer = '2026_001'").fetchone()[0]
    assert bezahlt == 0
