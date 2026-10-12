"""Übernimmt die bisherige Excel-Datei (Blätter AR_JJJJ und ER-JJJJ).

Aufbau der Excel-Datei:
* AR_JJJJ: Spalten ReNr, Empfänger, Datum, netto, brutto, bezahlt, MwSt, … Thema/Projekt
  (im Blatt AR_2020 ist alles um eine Spalte nach rechts verschoben).
* ER-JJJJ: Kopfzeile in Zeile 3, darunter zwölf Monatsblöcke, jeweils abgeschlossen
  durch eine Summenzeile `=SUM(Kx:Ky)`. Die Datumsspalte ist oft nicht gepflegt,
  daher wird der Monat aus dem Block abgeleitet.
"""
import re
from datetime import date

import openpyxl

from .geld import MONATE, euro_zu_cent
from .logik import ar_speichern, er_speichern, kunde_finden_oder_anlegen

TOLERANZ_CENT = 3
MONATSNAMEN = {m.lower() for m in MONATE} | {"januar"}
SUMMEN_FORMEL = re.compile(r"^=SUM\(K(\d+):K(\d+)\)$", re.I)


class ImportBericht:
    def __init__(self):
        self.ar = 0
        self.er = 0
        self.uebersprungen = 0
        self.warnungen = []

    def warnung(self, text):
        self.warnungen.append(text)

    def __str__(self):
        zeilen = [
            f"{self.ar} Ausgangsrechnungen und {self.er} Eingangsrechnungen übernommen, "
            f"{self.uebersprungen} bereits vorhandene übersprungen."
        ]
        if self.warnungen:
            zeilen.append(f"{len(self.warnungen)} Hinweise:")
            zeilen += [f"  - {w}" for w in self.warnungen]
        return "\n".join(zeilen)


def _text(wert):
    if wert is None or isinstance(wert, (int, float)):
        return ""
    text = " ".join(str(wert).split())
    if text.startswith("#") or text.lower() in MONATSNAMEN:
        return ""
    return text


def _datum(wert):
    if hasattr(wert, "date"):
        return wert.date()
    if isinstance(wert, date):
        return wert
    return None


def _cent(wert):
    try:
        return euro_zu_cent(wert) if isinstance(wert, (int, float)) else None
    except ValueError:
        return None


def importieren(con, pfad, offen_wert=1):
    """Importiert alle Jahre aus der Excel-Datei. Bereits vorhandene Rechnungen werden übersprungen.

    `offen_wert`: Wert der Spalte „bezahlt“, der eine *offene* Rechnung kennzeichnet.
    In der bisherigen Datei steht 1 bei den jüngsten, noch nicht bezahlten Rechnungen.
    """
    werte = openpyxl.load_workbook(pfad, data_only=True)
    formeln = openpyxl.load_workbook(pfad, data_only=False)
    bericht = ImportBericht()
    for name in werte.sheetnames:
        if m := re.fullmatch(r"AR_(\d{4})", name):
            _ar_blatt(con, werte[name], int(m.group(1)), offen_wert, bericht)
        elif m := re.fullmatch(r"ER-(\d{4})", name):
            _er_blatt(con, werte[name], formeln[name], int(m.group(1)), bericht)
    con.commit()
    return bericht


# --------------------------------------------------------------------- AR

def _ar_blatt(con, blatt, jahr, offen_wert, bericht):
    # Spaltenversatz ermitteln (Kopfzeile „ReNr“ in Spalte A oder B)
    versatz, kopfzeile = 0, 1
    for zeile in blatt.iter_rows(min_row=1, max_row=5):
        for zelle in zeile[:3]:
            if isinstance(zelle.value, str) and zelle.value.strip().lower() == "renr":
                versatz, kopfzeile = zelle.column - 1, zelle.row
    gesehen = set()
    for zeile in blatt.iter_rows(min_row=kopfzeile + 1, values_only=True):
        zeile = zeile[versatz:] + (None,) * 12
        nummer = zeile[0]
        if not (isinstance(nummer, str) and re.fullmatch(rf"{jahr}_\d+(_\d+)?", nummer.strip())):
            continue
        nummer = nummer.strip()
        empfaenger, datum, netto, brutto, bezahlt = zeile[1], _datum(zeile[2]), _cent(zeile[3]), _cent(zeile[4]), zeile[5]
        if not _text(empfaenger) or netto is None:
            continue  # leere Vorlagezeile
        if nummer in gesehen:
            bericht.warnung(f"{nummer}: doppelte Rechnungsnummer im Blatt, zweite Zeile ignoriert.")
            continue
        gesehen.add(nummer)
        if con.execute("SELECT 1 FROM ausgangsrechnungen WHERE nummer = ?", (nummer,)).fetchone():
            bericht.uebersprungen += 1
            continue
        if datum is None:
            datum = date(jahr, 1, 1)
            bericht.warnung(f"{nummer}: kein Datum, 01.01.{jahr} eingesetzt.")
        elif datum.year != jahr:
            bericht.warnung(f"{nummer}: Datum {datum:%d.%m.%Y} liegt nicht im Jahr {jahr} – bitte prüfen.")

        notizen = [_text(v) for v in zeile[7:10]]
        projekt = notizen[2] if blatt.cell(kopfzeile, versatz + 10).value else ""
        notiz = " / ".join(t for t in notizen if t and t != projekt)
        if brutto is None:
            brutto = 0

        positionen = _ar_positionen(nummer, netto, brutto, projekt or notiz, bericht)
        ar_speichern(
            con,
            {
                "datum": datum,
                "kunde_id": kunde_finden_oder_anlegen(con, empfaenger),
                "projekt": projekt,
                "notiz": notiz,
                "faellig_am": datum,
                "bezahlt": 0 if bezahlt == offen_wert else 1,
                "importiert": 1,
                "steuerhinweis": "Reverse Charge – die Steuerschuld geht auf den Leistungsempfänger über."
                if "reverse" in f"{projekt} {notiz}".lower() else "",
            },
            positionen,
            nummer=nummer,
        )
        bericht.ar += 1


def _ar_positionen(nummer, netto, brutto, text, bericht):
    """Leitet aus Netto/Brutto der Excel-Zeile Steuersatz und Positionen ab."""
    beschreibung = text or "Leistung laut Rechnung"
    if abs(brutto - netto) <= TOLERANZ_CENT:
        return [{"beschreibung": beschreibung, "einzelpreis_cent": netto, "ust_satz": 0}]
    if brutto == 0:
        bericht.warnung(f"{nummer}: kein Bruttobetrag in Excel – als steuerfrei (0 %) übernommen.")
        return [{"beschreibung": beschreibung, "einzelpreis_cent": netto, "ust_satz": 0}]
    for satz in (20, 10, 13):
        if abs(brutto - round(netto * (100 + satz) / 100)) <= TOLERANZ_CENT:
            return [{"beschreibung": beschreibung, "einzelpreis_cent": netto, "ust_satz": satz}]
    rest = brutto - round(netto * 1.2)
    if rest > 0:
        # Honorar + 20 % USt + Barauslagen ohne USt (z. B. Bahn, Hotel)
        bericht.warnung(f"{nummer}: Brutto enthält {rest / 100:.2f} € ohne USt – als Barauslagen-Position übernommen.")
        return [
            {"beschreibung": beschreibung, "einzelpreis_cent": netto, "ust_satz": 20},
            {"beschreibung": "Barauslagen / Spesen (ohne USt)", "einzelpreis_cent": rest, "ust_satz": 0},
        ]
    bericht.warnung(f"{nummer}: Brutto {brutto / 100:.2f} € passt nicht zu Netto {netto / 100:.2f} € – mit 20 % neu berechnet.")
    return [{"beschreibung": beschreibung, "einzelpreis_cent": netto, "ust_satz": 20}]


# --------------------------------------------------------------------- ER

def _er_bloecke(formelblatt):
    """Liefert die Zeilenbereiche der zwölf Monatsblöcke anhand der Summenformeln in Spalte K."""
    bloecke = []
    for zeile in formelblatt.iter_rows(min_row=4, min_col=11, max_col=11):
        wert = zeile[0].value
        if isinstance(wert, str) and (m := SUMMEN_FORMEL.match(wert.replace(" ", ""))):
            von, bis = int(m.group(1)), int(m.group(2))
            if bis - von < 120:  # Gesamtsummen über das ganze Jahr ignorieren
                bloecke.append((von, bis))
        if len(bloecke) == 12:
            break
    return bloecke


def _er_blatt(con, blatt, formelblatt, jahr, bericht):
    if con.execute("SELECT 1 FROM eingangsrechnungen WHERE jahr = ? AND importiert = 1", (jahr,)).fetchone():
        bericht.uebersprungen += 1
        bericht.warnung(f"ER-{jahr}: bereits importiert, Blatt übersprungen.")
        return
    bloecke = _er_bloecke(formelblatt)
    if len(bloecke) < 12:
        bericht.warnung(f"ER-{jahr}: nur {len(bloecke)} Monatsblöcke erkannt.")
    for monat, (von, bis) in enumerate(bloecke, start=1):
        for nr in range(von, bis + 1):
            z = [blatt.cell(nr, s).value for s in range(1, 20)]
            brutto = _cent(z[5])
            if not brutto:
                continue
            enr = z[0] if isinstance(z[0], int) else None
            lieferant = _text(z[3])
            text = " – ".join(t for t in (_text(z[7]), _text(z[8])) if t)
            # Fixkosten (Miete, Leasing, Versicherung) stehen in Excel oft ohne Lieferant
            lieferant = lieferant or text

            # Datum: nur übernehmen, wenn es zum Monatsblock passt
            d = _datum(z[1])
            geschaetzt = not (d and d.year == jahr and d.month == monat)
            datum = date(jahr, monat, 1) if geschaetzt else d

            # Steuersatz aus der berechneten MwSt-Spalte ableiten
            mwst = _cent(z[11]) or 0
            satz, manuell = 20, None
            for s in (20, 10, 13, 0):
                if abs(mwst - round(brutto * s / (100 + s))) <= TOLERANZ_CENT:
                    satz = s
                    break
            else:
                manuell = mwst

            ziel = _text(z[16])
            verrechnet = _text(z[17]).lower()
            ueberwiesen = _text(z[18]).lower()
            weiter_kunde = None
            if ziel and not ziel.lower().startswith("futura") and not verrechnet.startswith("intern"):
                weiter_kunde = kunde_finden_oder_anlegen(con, ziel)
            notiz = f"Excel: weiterverrechnen an {ziel}" if ziel and not weiter_kunde and not ziel.lower().startswith("futura") else ""

            er_speichern(con, {
                "datum": datum,
                "datum_geschaetzt": int(geschaetzt),
                "enr": enr,
                "lieferant": lieferant,
                "text": text,
                "kostenart": _text(z[9]),
                "brutto_cent": brutto,
                "ust_satz": satz,
                "vorsteuer_manuell_cent": manuell,
                "bezahlt": ueberwiesen not in ("nein", "???"),
                "weiter_kunde_id": weiter_kunde,
                "weiter_erledigt": verrechnet == "ja",
                "notiz": notiz,
                "importiert": 1,
            })
            bericht.er += 1
