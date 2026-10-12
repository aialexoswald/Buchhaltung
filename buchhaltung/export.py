"""Excel-Export eines Jahres (z. B. für die Steuerberatung)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EURO = '#,##0.00 "€"'


def _blatt(wb, titel, kopf, zeilen, geldspalten):
    ws = wb.create_sheet(titel)
    ws.append(kopf)
    for zelle in ws[1]:
        zelle.font = Font(bold=True)
    for z in zeilen:
        ws.append(z)
    for spalte in geldspalten:
        for zelle in ws.iter_cols(min_col=spalte, max_col=spalte, min_row=2):
            for c in zelle:
                c.number_format = EURO
    letzte = ws.max_row
    if letzte > 1:
        ws.append([])
        summe = ["Summe"] + [None] * (len(kopf) - 1)
        for spalte in geldspalten:
            buchstabe = get_column_letter(spalte)
            summe[spalte - 1] = f"=SUM({buchstabe}2:{buchstabe}{letzte})"
        ws.append(summe)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
            c.number_format = EURO
    for i, _ in enumerate(kopf, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "A2"


def jahr_exportieren(con, jahr):
    wb = Workbook()
    wb.remove(wb.active)
    ar = con.execute(
        """SELECT ar.*, k.name AS kunde FROM ausgangsrechnungen ar JOIN kunden k ON k.id = ar.kunde_id
           WHERE substr(ar.datum, 1, 4) = ? ORDER BY ar.nummer""",
        (str(jahr),),
    ).fetchall()
    _blatt(
        wb, f"AR_{jahr}",
        ["ReNr", "Kunde", "Datum", "Projekt", "Netto", "USt", "Brutto", "Fällig", "Bezahlt am", "Status", "Notiz"],
        [
            [r["nummer"], r["kunde"], r["datum"], r["projekt"], r["netto_cent"] / 100, r["ust_cent"] / 100,
             r["brutto_cent"] / 100, r["faellig_am"], r["bezahlt_am"],
             "storniert" if r["storniert"] else "Storno" if r["storno_von_id"] else "bezahlt" if r["bezahlt"] else "offen",
             r["notiz"]]
            for r in ar
        ],
        [5, 6, 7],
    )
    er = con.execute(
        """SELECT er.*, k.name AS weiter_an FROM eingangsrechnungen er
           LEFT JOIN kunden k ON k.id = er.weiter_kunde_id WHERE er.jahr = ? ORDER BY er.datum, er.enr""",
        (jahr,),
    ).fetchall()
    _blatt(
        wb, f"ER-{jahr}",
        ["ENR", "Datum", "Lieferant", "Re-Nr. Lieferant", "Text", "Kostenart", "Brutto", "USt-Satz", "Netto",
         "Vorsteuer", "Reverse Charge", "Bezahlt", "Weiterverrechnen an", "Weiterverrechnet", "Notiz"],
        [
            [r["enr"], r["datum"], r["lieferant"], r["rechnungsnr_lieferant"], r["text"], r["kostenart"],
             r["brutto_cent"] / 100, r["ust_satz"], r["netto_cent"] / 100, r["vorsteuer_cent"] / 100,
             "ja" if r["reverse_charge"] else "", "ja" if r["bezahlt"] else "nein", r["weiter_an"] or "intern",
             ("ja" if r["weiter_erledigt"] else "nein") if r["weiter_an"] else "", r["notiz"]]
            for r in er
        ],
        [7, 9, 10],
    )
    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()
